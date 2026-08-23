#!/usr/bin/env python3
"""Create deterministic, seeded entity variants of Harvey LAB tasks.

The Harvey mirror is always read-only. A generation is assembled in a sibling
staging directory, fully re-derived and validated, and only then published. An
existing seed is never replaced unless ``--replace-generated`` is explicit and
the existing manifest identifies the exact same source task and seed.

DOCX text is replaced within WordprocessingML text runs (including deleted
runs inside tracked changes) and, when necessary, across runs in the same
paragraph; a boundary-guarded pass over every serialized XML part also covers
tracked-change metadata (``w:author``), reviewer rosters, and docProps text.
Other OOXML formats are changed only by raw XML text substitution, preserving
package member order and ZIP metadata. A
semantic residual scan fails closed when a source entity is split or encoded in
a form the structure-preserving mutation cannot safely rewrite.

Usage:
    python3 research/lab_mutate.py \
        --task banking-finance/identify-issues-in-compliance-certificate \
        --entities research/mutation-configs/<task>/entities.json \
        --seed 1 --out research/mutations

    python3 research/lab_mutate.py \
        --check research/mutations/<area>/<task>__seed-001
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import fcntl
import hashlib
import json
import os
import random
import re
import shutil
import stat
import struct
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from email import policy
from email.parser import BytesParser
from pathlib import Path, PurePosixPath
from typing import Any, NamedTuple
from xml.etree import ElementTree


ROOT = Path(__file__).resolve().parents[1]
LAB = ROOT / "research" / "repos" / "harveyai@harvey-labs"
DEFAULT_OUTPUT = ROOT / "research" / "mutations"
DEFAULT_PLAN = ROOT / "research" / "mutation-configs" / "seed-plan.json"
DEFAULT_PIN_FILE = ROOT / "research" / "repos-commits.json"
PIN_KEY = "harveyai@harvey-labs"
SOURCE_REPO = "harveyai/harvey-labs"
TOOL_VERSION = 3
OOXML_EXTENSIONS = {".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"}
WORD_EXTENSIONS = {".docx", ".docm"}
PLAIN_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".xml", ".html", ".htm"}
XML_PART_EXTENSIONS = {".xml", ".rels"}

# Name banks for deterministic replacement generation. Synthetic only, per
# harvey-labs CONTRIBUTING.md ground rules.
COMPANY_HEADS = [
    "Juniper", "Blue Ridge", "Silverbirch", "Kestrel", "Alderpoint",
    "Stonegate", "Hollis", "Merribrook", "Larchmont", "Douglas Fork",
    "Timberline", "Oakhaven", "Redcedar", "Granite Bay", "Wolfcreek",
    "Bellhaven", "Corvid", "Eastlight", "Fenwick", "Halloran",
    "Ironvale", "Marrowgate", "Northquay", "Quillbrook", "Tarrowfield",
]
COMPANY_TAILS = {
    "company": ["Sawmills", "Forest Products", "Woodworks", "Timberworks", "Millcraft"],
    "holdings": ["Capital Holdings", "Timber Holdings", "Resource Holdings", "Industrial Holdings"],
    "partners": ["Partners", "Equity Partners", "Growth Partners", "Capital Partners"],
    "target": ["Lumber Co", "Plywood Co", "Veneer Co", "Panel Co"],
    "accounting": ["Accounting Group", "Audit Group"],
    "advisory": ["Advisory", "Advisors"],
    "products": ["Panel Products", "Board Products", "Timber Products"],
}
FIRST_NAMES_F = [
    "Alicia", "Priya", "Carmen", "Ingrid", "Elena", "Naomi", "Fatima",
    "Mireille", "Bettina", "Soraya", "Willa", "Damaris", "Katya", "Noor",
]
FIRST_NAMES_M = [
    "Marcus", "Theo", "Rashid", "Victor", "Owen", "Daniel", "Anders",
    "Emeka", "Lorcan", "Matteo", "Hugo", "Tobias", "Ravi", "Callum",
]
LAST_NAMES = [
    "Calloway", "Bram", "Osei", "Lindqvist", "Marchetti", "Duval",
    "Hartwell", "Nakata", "Reyes", "Kovacs", "Whitcombe", "Ferro",
    "Ashford", "Beaumont", "Castellanos", "Drummond", "Eriksen", "Fontaine",
    "Grantham", "Iwasaki", "Joubert", "Krishnan", "Mireles", "Okonkwo",
]
PLACES = ["Washington", "Idaho", "Montana", "Vermont", "Maine", "Georgia"]


def canonical_json(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_xml_part(name: str) -> bool:
    return name.lower().endswith((".xml", ".rels"))


def safe_relative(value: str, label: str) -> PurePosixPath:
    path = PurePosixPath(value)
    if (
        path.is_absolute()
        or not path.parts
        or "\\" in value
        or any(ord(character) < 32 for character in value)
        or any(part in {"", ".", ".."} for part in path.parts)
    ):
        raise ValueError(f"unsafe {label}: {value!r}")
    return path


def ensure_safe_descendant(path: Path, root: Path, label: str) -> Path:
    root = root.resolve()
    try:
        relative = path.relative_to(root)
        resolved = path.resolve(strict=False)
        resolved.relative_to(root)
    except ValueError as error:
        raise ValueError(f"{label} escapes its root: {path}") from error
    cursor = root
    for part in relative.parts:
        cursor /= part
        if cursor.is_symlink():
            raise ValueError(f"{label} traverses a symlink: {cursor}")
    return resolved


def git_output(source: Path, *arguments: str) -> str:
    return subprocess.run(
        ["git", "-C", str(source), *arguments],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()


def git_commit(source: Path) -> str:
    return git_output(source, "rev-parse", "HEAD")


def git_clean(source: Path) -> bool:
    return not bool(git_output(source, "status", "--porcelain", "--untracked-files=all"))


def git_remote(source: Path) -> str:
    return git_output(source, "remote", "get-url", "origin")


def normalized_remote(remote: str) -> str:
    value = remote.strip().rstrip("/")
    if value.endswith(".git"):
        value = value[:-4]
    if value.startswith("git@github.com:"):
        value = "https://github.com/" + value.removeprefix("git@github.com:")
    return value.lower()


def validate_source_identity(source: Path, commit: str, remote: str) -> str | None:
    """Bind the default mirror to its checked-in pin and canonical GitHub remote."""
    if source.resolve() != LAB.resolve():
        return None
    if normalized_remote(remote) != "https://github.com/harveyai/harvey-labs":
        raise ValueError(f"default Harvey mirror has unexpected origin: {remote}")
    pins = json.loads(DEFAULT_PIN_FILE.read_text(encoding="utf-8"))
    expected = pins.get(PIN_KEY)
    if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-fA-F]{7,40}", expected):
        raise ValueError(f"missing or invalid Harvey source pin in {DEFAULT_PIN_FILE}")
    if not commit.startswith(expected):
        raise ValueError(f"source commit mismatch: expected {expected}, got {commit}")
    return expected


class SourceAttestation(NamedTuple):
    root: Path
    commit: str
    remote: str
    expected_commit: str | None
    license_sha256: str


def attest_source(source_root: Path) -> SourceAttestation:
    """Capture one clean, pinned source identity for a bounded operation."""
    source_root = source_root.resolve()
    if not git_clean(source_root):
        raise ValueError("source Harvey worktree is dirty")
    commit = git_commit(source_root)
    remote = git_remote(source_root)
    expected_commit = validate_source_identity(source_root, commit, remote)
    source_license = source_root / "LICENSE"
    if not source_license.is_file() or source_license.is_symlink():
        raise ValueError("source Harvey license is missing or unsafe")
    return SourceAttestation(
        root=source_root,
        commit=commit,
        remote=remote,
        expected_commit=expected_commit,
        license_sha256=sha256_file(source_license),
    )


def use_source_attestation(
    source_root: Path, attestation: SourceAttestation | None
) -> tuple[SourceAttestation, bool]:
    source_root = source_root.resolve()
    if attestation is None:
        return attest_source(source_root), True
    if attestation.root != source_root:
        raise ValueError("source attestation belongs to a different mirror")
    return attestation, False


def verify_source_attestation(attestation: SourceAttestation) -> None:
    """Prove a shared plan/check attestation is still current at closeout."""
    source_root = attestation.root
    current_commit = git_commit(source_root)
    current_remote = git_remote(source_root)
    if current_commit != attestation.commit or current_remote != attestation.remote:
        raise ValueError("source Harvey identity changed during operation")
    if (
        validate_source_identity(source_root, current_commit, current_remote)
        != attestation.expected_commit
    ):
        raise ValueError("source Harvey pin changed during operation")
    source_license = source_root / "LICENSE"
    if (
        not source_license.is_file()
        or source_license.is_symlink()
        or sha256_file(source_license) != attestation.license_sha256
    ):
        raise ValueError("source Harvey license changed during operation")
    if not git_clean(source_root):
        raise ValueError("source Harvey worktree changed during operation")


def validate_seed(seed: Any) -> int:
    if isinstance(seed, bool) or not isinstance(seed, int) or seed < 0:
        raise ValueError("seed must be a non-negative integer")
    return seed


@contextmanager
def mutation_output_lock(output_root: Path):
    """Serialize generators and checkers that share one mutation output tree."""
    identity = hashlib.sha256(str(output_root.resolve()).encode()).hexdigest()[:24]
    lock_path = Path(tempfile.gettempdir()) / (
        f"legal-agent-mutations-{os.getuid()}-{identity}.lock"
    )
    flags = os.O_CREAT | os.O_RDWR
    if hasattr(os, "O_NOFOLLOW"):
        flags |= os.O_NOFOLLOW
    descriptor = os.open(lock_path, flags, 0o600)
    metadata = os.fstat(descriptor)
    if (
        not stat.S_ISREG(metadata.st_mode)
        or metadata.st_uid != os.getuid()
        or metadata.st_nlink != 1
    ):
        os.close(descriptor)
        raise RuntimeError(f"unsafe mutation lock file: {lock_path}")
    os.fchmod(descriptor, 0o600)
    with os.fdopen(descriptor, "a+", encoding="utf-8") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def kebab(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def squash(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def acronym_of(value: str) -> str:
    return "".join(word[0] for word in value.split() if word and word[0].isupper()).upper()


def validate_entity_text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} must be a non-empty string")
    if value != value.strip() or any(ord(character) < 32 for character in value):
        raise ValueError(f"{label} contains leading/trailing or control whitespace")
    if "/" in value or "\\" in value:
        raise ValueError(f"{label} contains a path separator")
    return value


def validate_entities(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list) or not value:
        raise ValueError("entities config must be a non-empty JSON array")
    names: set[str] = set()
    for index, entity in enumerate(value):
        if not isinstance(entity, dict):
            raise ValueError(f"entity {index} must be an object")
        entity_type = entity.get("type")
        if entity_type not in {"company", "acronym_company", "person", "place"}:
            raise ValueError(f"entity {index} has unknown type: {entity_type!r}")
        name = validate_entity_text(entity.get("name"), f"entity {index} name")
        if name in names:
            raise ValueError(f"duplicate entity name: {name!r}")
        names.add(name)

        fragments = entity.get("fragments", [])
        if not isinstance(fragments, list):
            raise ValueError(f"entity {index} fragments must be a list")
        for fragment_index, fragment in enumerate(fragments):
            validate_entity_text(fragment, f"entity {index} fragment {fragment_index}")

        if entity_type in {"company", "acronym_company"}:
            role = entity.get("role", "company")
            if role != "lawfirm" and role not in COMPANY_TAILS:
                raise ValueError(f"entity {index} has unknown company role: {role!r}")
            tails = entity.get("tail_options")
            if tails is not None:
                if not isinstance(tails, list) or not tails:
                    raise ValueError(f"entity {index} tail_options must be a non-empty list")
                for tail_index, tail in enumerate(tails):
                    validate_entity_text(tail, f"entity {index} tail option {tail_index}")
            if entity.get("domain") is not None:
                domain = validate_entity_text(entity["domain"], f"entity {index} domain")
                if not re.fullmatch(r"[A-Za-z0-9.-]+", domain):
                    raise ValueError(f"entity {index} domain is not a plain hostname")
            if entity_type == "acronym_company":
                acronym = validate_entity_text(entity.get("acronym"), f"entity {index} acronym")
                if not re.fullmatch(r"[A-Za-z0-9.-]+", acronym):
                    raise ValueError(f"entity {index} acronym contains unsupported punctuation")
        elif entity_type == "person":
            if len(name.split()) != 2:
                raise ValueError(f"person entity {name!r} must contain exactly first and last name")
            if entity.get("pronouns") not in {None, "she", "he"}:
                raise ValueError(f"person entity {name!r} pronouns must be 'she' or 'he'")
    return value


def _pick(rng: random.Random, bank: list[str], used: set[str]) -> str:
    choices = [entry for entry in bank if entry not in used]
    if not choices:
        raise ValueError("name bank exhausted; extend the deterministic name banks")
    choice = rng.choice(choices)
    used.add(choice)
    return choice


def _boundary_pattern(term: str) -> re.Pattern[str]:
    return re.compile(r"(?<![A-Za-z0-9])" + re.escape(term) + r"(?![A-Za-z0-9])")


def build_replacement_map(
    entities: list[dict[str, Any]],
    seed: int,
    corpus_text: str = "",
) -> list[tuple[str, str]]:
    """Return deterministic substitution pairs, longest source first."""
    validate_seed(seed)
    rng = random.Random(f"lab-mutate-v2:{seed}")
    corpus_words = set(re.findall(r"[a-z0-9]+", corpus_text.lower()))

    def clean(bank: list[str]) -> list[str]:
        return [
            entry for entry in bank
            if not any(word.lower() in corpus_words for word in entry.split() if len(word) > 3)
        ]

    heads, firsts_f, firsts_m, lasts, places = (
        clean(COMPANY_HEADS),
        clean(FIRST_NAMES_F),
        clean(FIRST_NAMES_M),
        clean(LAST_NAMES),
        clean(PLACES),
    )
    used: set[str] = set()
    pairs: list[tuple[str, str]] = []

    for entity in entities:
        entity_type = entity["type"]
        old = entity["name"]
        if entity_type in {"company", "acronym_company"}:
            role = entity.get("role", "company")
            if role == "lawfirm":
                new = f"{_pick(rng, lasts, used)} {_pick(rng, lasts, used)}"
            else:
                head = _pick(rng, heads, used)
                bank = entity.get("tail_options") or COMPANY_TAILS[role]
                tails = [tail for tail in bank if tail.split()[0].lower() != head.split()[-1].lower()]
                new = f"{head} {rng.choice(tails or bank)}"
            pairs.append((old, new))
            old_words, new_words = old.split(), new.split()
            for length in range(2, len(old_words)):
                keep = len(new_words) - (len(old_words) - length)
                pairs.append((" ".join(old_words[:length]), " ".join(new_words[:max(1, keep)])))
            for fragment in entity.get("fragments", []):
                if fragment == old:
                    continue
                fragment_words = fragment.split()
                if old_words[:len(fragment_words)] == fragment_words:
                    keep = len(new_words) - (len(old_words) - len(fragment_words))
                    new_fragment = " ".join(new_words[:max(1, keep)])
                elif old_words[-len(fragment_words):] == fragment_words:
                    keep = len(new_words) - (len(old_words) - len(fragment_words))
                    new_fragment = " ".join(new_words[-max(1, keep):])
                else:
                    raise ValueError(f"fragment {fragment!r} is not a prefix/suffix of {old!r}")
                pairs.append((fragment, new_fragment))
            if entity_type == "acronym_company":
                pairs.append((entity["acronym"], acronym_of(new)))
            if entity.get("domain"):
                pairs.append((entity["domain"], squash(new) + ".com"))
        elif entity_type == "person":
            first, last = old.split()
            pronouns = entity.get("pronouns")
            bank = firsts_f if pronouns == "she" else firsts_m if pronouns == "he" else firsts_f + firsts_m
            new_first = _pick(rng, bank, used)
            new_last = _pick(rng, lasts, used)
            pairs.extend([
                (old, f"{new_first} {new_last}"),
                (first, new_first),
                (last, new_last),
                ((first[0] + last).lower(), (new_first[0] + new_last).lower()),
            ])
        else:
            pairs.append((old, _pick(rng, places, used)))

    expanded: dict[str, str] = {}
    expanded_origins: dict[str, str] = {}
    for old, new in pairs:
        for source, destination in (
            (old, new),
            (old.lower(), new.lower()),
            (old.upper(), new.upper()),
            (kebab(old), kebab(new)),
            (squash(old), squash(new)),
        ):
            if not source or not destination or source == destination:
                continue
            previous = expanded.get(source)
            if (
                previous is not None
                and previous != destination
                and expanded_origins[source] != old
            ):
                raise ValueError(
                    f"ambiguous generated substitution for {source!r}: {previous!r} vs {destination!r}"
                )
            if previous is None:
                expanded[source] = destination
                expanded_origins[source] = old

    ordered = sorted(expanded.items(), key=lambda item: (-len(item[0]), item[0]))
    for old, _destination in ordered:
        for _source, new in ordered:
            if _boundary_pattern(old).search(new):
                raise ValueError(f"replacement output {new!r} still contains source term {old!r}")
    return ordered


def replace_text_count(
    text: str,
    pairs: list[tuple[str, str]],
) -> tuple[str, Counter[str]]:
    counts: Counter[str] = Counter()
    for old, new in pairs:
        text, occurrences = _boundary_pattern(old).subn(lambda _match, value=new: value, text)
        if occurrences:
            counts[old] += occurrences
    return text, counts


def replace_text(text: str, pairs: list[tuple[str, str]]) -> str:
    return replace_text_count(text, pairs)[0]


def deep_replace(value: Any, pairs: list[tuple[str, str]], counts: Counter[str]) -> Any:
    if isinstance(value, str):
        replaced, local_counts = replace_text_count(value, pairs)
        counts.update(local_counts)
        return replaced
    if isinstance(value, list):
        return [deep_replace(item, pairs, counts) for item in value]
    if isinstance(value, dict):
        return {key: deep_replace(item, pairs, counts) for key, item in value.items()}
    return value


WT_RE = re.compile(
    r"(<w:(?:t|delText)(?:\s[^>]*)?>)(.*?)(</w:(?:t|delText)>)", re.DOTALL
)
WP_RE = re.compile(r"<w:p(?:\s[^>]*)?>.*?</w:p>", re.DOTALL)


def xml_escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def xml_unescape(value: str) -> str:
    return (
        value.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
        .replace("&amp;", "&")
    )


def replace_across_runs(
    run_texts: list[str], pairs: list[tuple[str, str]]
) -> list[str]:
    """Replace paragraph text while retaining every original Word text run."""
    values = list(run_texts)
    for old, new in pairs:
        joined = "".join(values)
        matches = list(_boundary_pattern(old).finditer(joined))
        for match in reversed(matches):
            start, end = match.span()
            starts: list[int] = []
            ends: list[int] = []
            cursor = 0
            for value in values:
                starts.append(cursor)
                cursor += len(value)
                ends.append(cursor)
            first = next(index for index, run_end in enumerate(ends) if start < run_end)
            last = next(
                index
                for index, (run_start, run_end) in enumerate(zip(starts, ends))
                if end > run_start and end <= run_end
            )
            if first == last:
                local_start = start - starts[first]
                local_end = end - starts[first]
                values[first] = (
                    values[first][:local_start] + new + values[first][local_end:]
                )
                continue

            overlap_lengths = [
                max(0, min(end, ends[index]) - max(start, starts[index]))
                for index in range(first, last + 1)
            ]
            total_overlap = sum(overlap_lengths)
            boundaries = [0]
            cumulative = 0
            for length in overlap_lengths[:-1]:
                cumulative += length
                boundaries.append(len(new) * cumulative // total_overlap)
            boundaries.append(len(new))
            allocations = [
                new[boundaries[index]:boundaries[index + 1]]
                for index in range(len(overlap_lengths))
            ]
            first_prefix = values[first][:start - starts[first]]
            last_suffix = values[last][end - starts[last]:]
            for offset, index in enumerate(range(first, last + 1)):
                values[index] = allocations[offset]
            values[first] = first_prefix + values[first]
            values[last] += last_suffix
    return values


def render_word_text_run(match: re.Match[str], value: str) -> str:
    opening = match.group(1)
    if value and (value[0].isspace() or value[-1].isspace()) and "xml:space" not in opening:
        opening = opening[:-1] + ' xml:space="preserve">'
    return opening + xml_escape(value) + match.group(3)


def mutate_docx_xml(xml: str, pairs: list[tuple[str, str]]) -> str:
    def replace_paragraph(match: re.Match[str]) -> str:
        paragraph = match.group(0)
        runs = list(WT_RE.finditer(paragraph))
        if not runs:
            return paragraph
        originals = [xml_unescape(run.group(2)) for run in runs]
        replacements = replace_across_runs(originals, pairs)
        if replacements == originals:
            return paragraph
        replacement_iterator = iter(zip(originals, replacements))

        def render_replacement(run: re.Match[str]) -> str:
            original, replacement = next(replacement_iterator)
            if replacement == original:
                return run.group(0)
            return render_word_text_run(run, replacement)

        return WT_RE.sub(
            render_replacement,
            paragraph,
        )

    xml = WP_RE.sub(replace_paragraph, xml)

    def replace_unscoped_run(run: re.Match[str]) -> str:
        original = xml_unescape(run.group(2))
        replacement = replace_text(original, pairs)
        if replacement == original:
            return run.group(0)
        return render_word_text_run(run, replacement)

    return WT_RE.sub(replace_unscoped_run, xml)


def package_topology(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        rows = [
            {
                "name": info.filename,
                "compression": info.compress_type,
                "date_time": info.date_time,
                "flag_bits": info.flag_bits,
                "create_system": info.create_system,
                "create_version": info.create_version,
                "extract_version": info.extract_version,
                "external_attr": info.external_attr,
                "internal_attr": info.internal_attr,
                "extra": info.extra.hex(),
                "comment": info.comment.hex(),
                "directory": info.is_dir(),
            }
            for info in archive.infolist()
        ]
        archive_comment = archive.comment.hex()
    return {
        "parts": len(rows),
        "archive_comment": archive_comment,
        "sha256": sha256_bytes(canonical_json(rows)),
    }


def validate_ooxml(path: Path) -> int:
    xml_parts = 0
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"ZIP CRC failure in {path.name}:{bad_member}")
        for name in archive.namelist():
            if is_xml_part(name):
                ElementTree.fromstring(archive.read(name))
                xml_parts += 1
    extension = path.suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        import openpyxl

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_vba=extension == ".xlsm",
        )
        try:
            if not workbook.sheetnames:
                raise ValueError(f"workbook has no worksheets: {path.name}")
        finally:
            workbook.close()
    elif extension in {".docx", ".docm"}:
        from docx import Document

        Document(path)
    elif extension in {".pptx", ".pptm"}:
        from pptx import Presentation

        Presentation(path)
    return xml_parts


def replace_escaped_xml_boundary(xml: str, pairs: list[tuple[str, str]]) -> str:
    """Boundary-guarded substitution over serialized XML (attribute values,
    tracked-change metadata such as ``w:author``, deleted-run text, and
    docProps text nodes). Terms are XML-escaped before matching so the pass
    operates on exactly the byte forms the semantic residual scan reads; the
    boundary guards mean any document that previously passed the residual
    scan re-derives byte-identically under this pass."""
    for old, new in pairs:
        escaped_old, escaped_new = xml_escape(old), xml_escape(new)
        pattern = re.compile(
            r"(?<![A-Za-z0-9])" + re.escape(escaped_old) + r"(?![A-Za-z0-9])"
        )
        xml = pattern.sub(lambda _match, value=escaped_new: value, xml)
    return xml


DOCX_CONTENT_PART_RE = re.compile(
    r"word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml"
)

FAIL_CLAUSE_RE = re.compile(r"\bFAIL (?:only )?if\b")

# Local-file-header offset of the general-purpose bit flag (2 bytes, LE).
_LOCAL_HEADER_FLAG_OFFSET = 6
# Bit 3: sizes/CRC deferred to a data descriptor. Restoring a mismatch on this
# bit would desynchronize the member layout, so it fails closed instead.
_DATA_DESCRIPTOR_FLAG = 0x08


def snapshot_zip_member_metadata(
    archive: zipfile.ZipFile,
) -> list[tuple[str, int, int]]:
    """Capture (flag_bits, external_attr) per member BEFORE any write.

    ``ZipFile.writestr`` mutates the very ZipInfo objects handed to it, so the
    restoration pass must compare against copied values, not the objects.
    """
    # Preserve archive order instead of keying by filename: ZIP permits
    # duplicate member names, and collapsing those entries would restore the
    # wrong metadata to one of them.
    return [
        (info.filename, info.flag_bits, info.external_attr)
        for info in archive.infolist()
    ]


def restore_zip_member_metadata(
    output_archive: zipfile.ZipFile, source_infos: list[tuple[str, int, int]]
) -> None:
    """Round-trip ZIP member metadata that Python's writer normalizes.

    ``ZipFile.writestr`` resets each member's general-purpose flag bits and
    promotes a zero ``external_attr`` to ``0o600 << 16``; most upstream
    documents already carry exactly those normalized values, but one known
    upstream package (the white-collar negotiation memo) was authored by a
    different pipeline with ``flag_bits=0x802`` and ``external_attr=0``, so an
    otherwise byte-faithful copy failed the package-topology equality check.
    This pass restores the source values: ``external_attr`` lives only in the
    central directory (rewritten from these ZipInfo objects at close), while
    ``flag_bits`` is additionally patched into the already-written local
    header so the archive stays internally consistent. Members whose metadata
    already matches are untouched, which keeps every previously passing
    document re-deriving byte-identically.
    """
    if len(output_archive.filelist) != len(source_infos):
        raise ValueError("ZIP member count changed before metadata restoration")
    stream = output_archive.fp
    if stream is None:
        raise ValueError("ZIP output stream is closed before metadata restoration")
    for zinfo, (source_name, source_flag_bits, source_external_attr) in zip(
        output_archive.filelist, source_infos, strict=True
    ):
        if zinfo.filename != source_name:
            raise ValueError(
                "ZIP member order changed before metadata restoration: "
                f"expected {source_name!r}, got {zinfo.filename!r}"
            )
        if zinfo.external_attr != source_external_attr:
            zinfo.external_attr = source_external_attr
        if zinfo.flag_bits != source_flag_bits:
            if (zinfo.flag_bits ^ source_flag_bits) & _DATA_DESCRIPTOR_FLAG:
                raise ValueError(
                    f"cannot restore data-descriptor flag for {zinfo.filename}"
                )
            position = stream.tell()
            stream.seek(zinfo.header_offset + _LOCAL_HEADER_FLAG_OFFSET)
            stream.write(struct.pack("<H", source_flag_bits))
            stream.seek(position)
            zinfo.flag_bits = source_flag_bits


def mutate_docx(source: Path, destination: Path, pairs: list[tuple[str, str]]) -> None:
    with zipfile.ZipFile(source) as input_archive, zipfile.ZipFile(destination, "w") as output_archive:
        output_archive.comment = input_archive.comment
        source_metadata = snapshot_zip_member_metadata(input_archive)
        for info in input_archive.infolist():
            payload = input_archive.read(info.filename)
            if DOCX_CONTENT_PART_RE.fullmatch(info.filename):
                xml = mutate_docx_xml(payload.decode("utf-8"), pairs)
                payload = replace_escaped_xml_boundary(xml, pairs).encode("utf-8")
            elif is_xml_part(info.filename):
                payload = replace_escaped_xml_boundary(
                    payload.decode("utf-8"), pairs
                ).encode("utf-8")
            output_archive.writestr(info, payload)
        restore_zip_member_metadata(output_archive, source_metadata)
    if package_topology(source) != package_topology(destination):
        raise ValueError(f"DOCX package topology changed: {source.name}")
    validate_ooxml(destination)


def mutate_ooxml_raw(source: Path, destination: Path, pairs: list[tuple[str, str]]) -> None:
    with zipfile.ZipFile(source) as input_archive, zipfile.ZipFile(destination, "w") as output_archive:
        output_archive.comment = input_archive.comment
        source_metadata = snapshot_zip_member_metadata(input_archive)
        for info in input_archive.infolist():
            payload = input_archive.read(info.filename)
            if is_xml_part(info.filename):
                for old, new in pairs:
                    payload = payload.replace(
                        xml_escape(old).encode("utf-8"),
                        xml_escape(new).encode("utf-8"),
                    )
            output_archive.writestr(info, payload)
        restore_zip_member_metadata(output_archive, source_metadata)
    if package_topology(source) != package_topology(destination):
        raise ValueError(f"OOXML package topology changed: {source.name}")
    validate_ooxml(destination)


def mutate_plain(source: Path, destination: Path, pairs: list[tuple[str, str]]) -> None:
    text = source.read_text(encoding="utf-8")
    destination.write_text(replace_text(text, pairs), encoding="utf-8", newline="")
    if destination.suffix.lower() == ".json":
        json.loads(destination.read_text(encoding="utf-8"))


def mutate_eml(source: Path, destination: Path, pairs: list[tuple[str, str]]) -> None:
    message = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
    changed = False
    source_hash = sha256_file(source)
    for part_index, part in enumerate(message.walk()):
        headers = getattr(part, "_headers", [])
        for index, (name, value) in enumerate(headers):
            original = str(value)
            replaced = replace_text(original, pairs)
            if replaced != original:
                headers[index] = (name, replaced)
                changed = True
        if part.is_multipart() and not part.get_boundary():
            part.set_boundary(f"harvey-mutation-{source_hash[:24]}-{part_index}")
            changed = True
        if not part.is_multipart() and part.get_content_maintype() == "text":
            content = part.get_content()
            if isinstance(content, str):
                replaced = replace_text(content, pairs)
                if replaced != content:
                    part.set_content(replaced, subtype=part.get_content_subtype())
                    changed = True
    if changed:
        destination.write_bytes(message.as_bytes(policy=policy.default.clone(linesep="\n")))
    else:
        shutil.copy2(source, destination)


def ooxml_semantic_text(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not is_xml_part(name):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for element in root.iter():
                chunks.extend(str(value) for value in element.attrib.values())
                if element.text:
                    chunks.append(element.text)
                local_name = element.tag.rsplit("}", 1)[-1]
                if local_name in {"p", "si", "is", "txBody"}:
                    joined = "".join(element.itertext())
                    if joined:
                        chunks.append(joined)
    return "\n".join(chunks)


def extract_all_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension in OOXML_EXTENSIONS:
        return ooxml_semantic_text(path)
    if extension == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        chunks = [f"{name}: {value}" for name, value in message.items()]
        for part in message.walk():
            if not part.is_multipart() and part.get_content_maintype() == "text":
                content = part.get_content()
                if isinstance(content, str):
                    chunks.append(content)
        return "\n".join(chunks)
    return path.read_text(encoding="utf-8")


def validate_task(task: dict[str, Any]) -> None:
    if not isinstance(task, dict):
        raise ValueError("generated task must be a JSON object")
    for field in ("title", "instructions"):
        if not isinstance(task.get(field), str) or not task[field].strip():
            raise ValueError(f"generated task requires non-empty {field}")
    deliverables = task.get("deliverables")
    if not isinstance(deliverables, dict) or not deliverables:
        raise ValueError("generated task requires a non-empty deliverables mapping")
    deliverable_names: set[str] = set()
    for key, value in deliverables.items():
        if not isinstance(key, str) or not key.strip():
            raise ValueError("generated deliverable keys must be non-empty strings")
        safe_relative(key, "deliverable key")
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"generated deliverable {key!r} must map to a filename")
        safe_relative(value, f"deliverable {key!r}")
        deliverable_names.update({key, value})

    criteria = task.get("criteria")
    if not isinstance(criteria, list) or not criteria:
        raise ValueError("generated task requires non-empty criteria")
    identifiers: set[str] = set()
    for index, criterion in enumerate(criteria):
        if not isinstance(criterion, dict):
            raise ValueError(f"criterion {index} must be an object")
        for field in ("id", "title", "match_criteria"):
            if not isinstance(criterion.get(field), str) or not criterion[field].strip():
                raise ValueError(f"criterion {index} requires non-empty {field}")
        if criterion["id"] in identifiers:
            raise ValueError(f"duplicate criterion ID: {criterion['id']}")
        identifiers.add(criterion["id"])
        criterion_deliverables = criterion.get("deliverables")
        if not isinstance(criterion_deliverables, list) or not criterion_deliverables:
            raise ValueError(f"criterion {criterion['id']} requires a deliverables list")
        if any(
            not isinstance(value, str) or value not in deliverable_names
            for value in criterion_deliverables
        ):
            raise ValueError(f"criterion {criterion['id']} references an unknown deliverable")
        match_criteria = criterion["match_criteria"]
        # "FAIL if" and "FAIL only if" are both explicit fail triggers; the
        # upstream corpus uses the "only if" form (e.g. litigation
        # motion-to-dismiss C-021), which is fail-closed in the same sense —
        # a rubric with no FAIL clause is still rejected.
        if not match_criteria.startswith("PASS if") or not FAIL_CLAUSE_RE.search(match_criteria):
            raise ValueError(f"criterion {criterion['id']} requires fail-closed PASS/FAIL text")


def source_documents(task_dir: Path) -> list[Path]:
    documents = task_dir / "documents"
    if not documents.is_dir() or documents.is_symlink():
        raise ValueError(f"task-local documents directory is missing or unsafe: {documents}")
    if any(path.is_symlink() for path in documents.rglob("*")):
        raise ValueError(f"source documents tree contains a symlink: {documents}")
    paths = sorted(
        (path for path in documents.rglob("*") if path.is_file()),
        key=lambda path: path.relative_to(documents).as_posix(),
    )
    if not paths:
        raise ValueError("source task contains no documents")
    for path in paths:
        if path.is_symlink():
            raise ValueError(f"source document is a symlink: {path.relative_to(documents)}")
        if path.suffix.lower() not in OOXML_EXTENSIONS | PLAIN_EXTENSIONS | {".eml"}:
            raise ValueError(f"unsupported document format for fail-closed mutation: {path.name}")
    return paths


def mutated_relative_path(relative: Path, pairs: list[tuple[str, str]]) -> PurePosixPath:
    parts = [replace_text(part, pairs) for part in relative.parts]
    for part in parts:
        if "/" in part or "\\" in part:
            raise ValueError(f"mutation produced a path separator in {part!r}")
    return safe_relative("/".join(parts), "mutated document path")


def mutate_document(source: Path, destination: Path, pairs: list[tuple[str, str]]) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    extension = source.suffix.lower()
    if extension in WORD_EXTENSIONS:
        mutate_docx(source, destination, pairs)
    elif extension in OOXML_EXTENSIONS:
        mutate_ooxml_raw(source, destination, pairs)
    elif extension == ".eml":
        mutate_eml(source, destination, pairs)
    elif extension in PLAIN_EXTENSIONS:
        mutate_plain(source, destination, pairs)
    else:
        raise ValueError(f"unsupported document format: {source.name}")


def residual_terms(text: str, pairs: list[tuple[str, str]]) -> list[str]:
    found = []
    for term in sorted({old for old, _ in pairs if len(old) >= 4}, key=lambda value: (-len(value), value)):
        if re.search(
            r"(?<![A-Za-z0-9])" + re.escape(term) + r"(?![A-Za-z0-9])",
            text,
            flags=re.IGNORECASE,
        ):
            found.append(term)
    return found


def build_expected_task(
    source_task: dict[str, Any],
    pairs: list[tuple[str, str]],
    source_task_name: str,
    seed: int,
    source_commit: str,
    source_expected_commit: str | None,
    source_license_sha256: str,
    entities_sha256: str,
) -> tuple[dict[str, Any], Counter[str]]:
    counts: Counter[str] = Counter()
    task = deep_replace(source_task, pairs, counts)
    task["provenance"] = {
        "mutated_from": source_task_name,
        "mutation_seed": seed,
        "mutation_tool": f"research/lab_mutate.py v{TOOL_VERSION}",
        "source_repo": SOURCE_REPO,
        "source_commit": source_commit,
        "source_expected_commit": source_expected_commit,
        "source_license": "MIT",
        "source_license_sha256": source_license_sha256,
        "entities_sha256": entities_sha256,
        "originals_immutable": True,
    }
    validate_task(task)
    return task, counts


def output_relative_for(task_name: str, seed: int) -> PurePosixPath:
    validate_seed(seed)
    task_relative = safe_relative(task_name, "task")
    return PurePosixPath(
        *task_relative.parts[:-1],
        f"{task_relative.parts[-1]}__seed-{seed:03d}",
    )


def validate_existing_target(target: Path, task: str, seed: int) -> None:
    validate_seed(seed)
    if target.is_symlink():
        raise ValueError(f"refusing to replace symlink output: {target}")
    manifest_path = target / "mutation-manifest.json"
    if not manifest_path.is_file() or manifest_path.is_symlink():
        raise ValueError(f"refusing to replace unrecognized output directory: {target}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest_seed = manifest.get("seed")
    if (
        manifest.get("task") != task
        or isinstance(manifest_seed, bool)
        or manifest_seed != seed
    ):
        raise ValueError(f"existing output manifest does not identify {task!r} seed {seed}")


def check_generated(
    task_dir: Path,
    source_root: Path = LAB,
    quiet: bool = False,
    source_attestation: SourceAttestation | None = None,
) -> None:
    if task_dir.is_symlink():
        raise ValueError(f"generated task directory is a symlink: {task_dir}")
    task_dir = task_dir.resolve()
    source_root = source_root.resolve()
    source_attestation, owns_attestation = use_source_attestation(
        source_root, source_attestation
    )
    manifest_path = task_dir / "mutation-manifest.json"
    task_path = task_dir / "task.json"
    entities_path = task_dir / "mutation-entities.json"
    for path in (manifest_path, task_path, entities_path):
        if not path.is_file() or path.is_symlink():
            raise ValueError(f"missing or unsafe generated artifact: {path}")

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("schema_version") != 2:
        raise ValueError("generated mutation manifest schema_version must be 2")
    source_task_name = manifest.get("task")
    seed = manifest.get("seed")
    if (
        not isinstance(source_task_name, str)
        or isinstance(seed, bool)
        or not isinstance(seed, int)
        or seed < 0
    ):
        raise ValueError("mutation manifest requires task and integer seed")
    task_relative = safe_relative(source_task_name, "manifest task")
    expected_output_task = output_relative_for(source_task_name, seed).as_posix()
    expected_manifest_values = {
        "tool": f"research/lab_mutate.py v{TOOL_VERSION}",
        "output_task": expected_output_task,
        "source_repo": SOURCE_REPO,
        "source_license": "MIT",
        "source_license_path": "LICENSE",
        "source_git_clean": True,
        "entity_count": None,
        "validation": {
            "source_entity_residuals": 0,
            "document_count_preserved": True,
            "source_immutable": True,
        },
    }
    for field, expected_value in expected_manifest_values.items():
        if field == "entity_count":
            continue
        if manifest.get(field) != expected_value:
            raise ValueError(f"mutation manifest has invalid {field}")

    current_commit = source_attestation.commit
    if current_commit != manifest.get("source_commit"):
        raise ValueError(
            f"source mirror moved: expected {manifest.get('source_commit')}, got {current_commit}"
        )
    current_remote = source_attestation.remote
    if current_remote != manifest.get("source_remote"):
        raise ValueError("source mirror remote differs from manifest")
    expected_commit = source_attestation.expected_commit
    if manifest.get("source_expected_commit") != expected_commit:
        raise ValueError("source pin differs from mutation manifest")
    source_license = source_root / "LICENSE"
    if (
        not source_license.is_file()
        or source_license.is_symlink()
        or sha256_file(source_license) != manifest.get("source_license_sha256")
    ):
        raise ValueError("source Harvey license differs from mutation manifest")

    source_task_dir = source_root / "tasks" / Path(*task_relative.parts)
    ensure_safe_descendant(
        source_task_dir,
        (source_root / "tasks").resolve(),
        "manifest source task",
    )
    source_task_path = source_task_dir / "task.json"
    if not source_task_path.is_file() or source_task_path.is_symlink():
        raise ValueError("source task.json is missing or unsafe")
    if sha256_file(source_task_path) != manifest.get("source_task_json_sha256"):
        raise ValueError("source task.json differs from manifest")

    entities = validate_entities(json.loads(entities_path.read_text(encoding="utf-8")))
    entities_sha256 = sha256_bytes(canonical_json(entities))
    if entities_sha256 != manifest.get("entities_sha256"):
        raise ValueError("mutation entities differ from manifest")
    if manifest.get("entity_count") != len(entities):
        raise ValueError("mutation manifest entity_count is incorrect")
    documents = source_documents(source_task_dir)
    source_docs_root = source_task_dir / "documents"
    corpus = "\n".join(
        [source_task_path.read_text(encoding="utf-8")]
        + [extract_all_text(path) for path in documents]
    )
    pairs = build_replacement_map(entities, seed, corpus)
    manifest_pairs = [(row["old"], row["new"]) for row in manifest.get("substitutions", [])]
    if pairs != manifest_pairs:
        raise ValueError("substitutions do not re-derive from entities and seed")

    source_task = json.loads(source_task_path.read_text(encoding="utf-8"))
    expected_task, task_counts = build_expected_task(
        source_task,
        pairs,
        source_task_name,
        seed,
        current_commit,
        expected_commit,
        sha256_file(source_license),
        entities_sha256,
    )
    if canonical_json(expected_task) != task_path.read_bytes():
        raise ValueError("generated task.json contains a non-derived change")
    if sha256_file(task_path) != manifest.get("output_task_json_sha256"):
        raise ValueError("generated task.json hash differs from manifest")
    if dict(sorted(task_counts.items())) != manifest.get("task_replacement_counts"):
        raise ValueError("task replacement counts differ from manifest")

    rows = manifest.get("files")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("mutation manifest files must be a list")
    source_rows = {row.get("source_path"): row for row in rows}
    output_rows = {row.get("output_path"): row for row in rows}
    if len(source_rows) != len(rows) or len(output_rows) != len(rows):
        raise ValueError("mutation manifest contains duplicate source/output paths")
    current_source_paths = {
        path.relative_to(source_docs_root).as_posix(): path for path in documents
    }
    if set(current_source_paths) != set(source_rows):
        raise ValueError("source document paths differ from generation manifest")
    output_docs = task_dir / "documents"
    if not output_docs.is_dir() or output_docs.is_symlink():
        raise ValueError("generated documents directory is missing or unsafe")
    actual_output_paths = {
        path.relative_to(output_docs).as_posix(): path
        for path in output_docs.rglob("*")
        if path.is_file()
    }
    if any(path.is_symlink() for path in task_dir.rglob("*")):
        raise ValueError("generated documents contain a symlink")
    if set(actual_output_paths) != set(output_rows):
        raise ValueError("generated document paths differ from manifest")

    with tempfile.TemporaryDirectory(prefix="harvey-mutation-check-") as temporary:
        temporary_root = Path(temporary)
        for source_relative, source_path in sorted(current_source_paths.items()):
            row = source_rows[source_relative]
            output_relative = mutated_relative_path(Path(source_relative), pairs).as_posix()
            if output_relative != row.get("output_path"):
                raise ValueError(f"output path does not re-derive: {source_relative}")
            output_path = actual_output_paths[output_relative]
            if sha256_file(source_path) != row.get("source_sha256"):
                raise ValueError(f"source document differs from manifest: {source_relative}")
            if sha256_file(output_path) != row.get("output_sha256"):
                raise ValueError(f"generated document differs from manifest: {output_relative}")
            if row.get("extension") != source_path.suffix.lower():
                raise ValueError(f"manifest extension is incorrect: {source_relative}")
            if row.get("source_bytes") != source_path.stat().st_size:
                raise ValueError(f"manifest source byte count is incorrect: {source_relative}")
            if row.get("output_bytes") != output_path.stat().st_size:
                raise ValueError(f"manifest output byte count is incorrect: {output_relative}")
            expected_path = temporary_root / output_relative
            mutate_document(source_path, expected_path, pairs)
            if expected_path.read_bytes() != output_path.read_bytes():
                raise ValueError(f"generated document contains a non-derived change: {output_relative}")
            if output_path.suffix.lower() in OOXML_EXTENSIONS:
                if package_topology(source_path) != package_topology(output_path):
                    raise ValueError(f"OOXML package topology differs: {output_relative}")
                if row.get("package_topology_preserved") is not True:
                    raise ValueError(f"manifest does not confirm OOXML topology: {output_relative}")
                validate_ooxml(source_path)
                validate_ooxml(output_path)
            elif row.get("package_topology_preserved") is not None:
                raise ValueError(f"manifest has unexpected package topology value: {output_relative}")
            terms = residual_terms(extract_all_text(output_path), pairs)
            terms.extend(residual_terms(output_relative, pairs))
            if terms:
                raise ValueError(f"residual source terms in {output_relative}: {sorted(set(terms))}")

    expected_files = {
        "task.json",
        "mutation-entities.json",
        "mutation-manifest.json",
        *{f"documents/{path}" for path in output_rows},
    }
    actual_files = {
        path.relative_to(task_dir).as_posix()
        for path in task_dir.rglob("*")
        if path.is_file()
    }
    if expected_files != actual_files:
        raise ValueError("generated tree contains missing or unexpected files")
    if residual_terms(task_path.read_text(encoding="utf-8"), pairs):
        raise ValueError("generated task.json contains residual source terms")
    if owns_attestation:
        verify_source_attestation(source_attestation)
    if not quiet:
        print(
            f"seeded Harvey task valid: {source_task_name} seed {seed:03d} "
            f"({len(rows)} documents, source {current_commit[:12]})"
        )


def generate(
    task_name: str,
    entities_path: Path,
    seed: int,
    source_root: Path,
    output_root: Path,
    replace_generated: bool = False,
    source_attestation: SourceAttestation | None = None,
) -> Path:
    validate_seed(seed)
    task_relative = safe_relative(task_name, "task")
    source_root = source_root.resolve()
    output_root = output_root.resolve()
    source_attestation, owns_attestation = use_source_attestation(
        source_root, source_attestation
    )
    source_commit = source_attestation.commit
    source_remote = source_attestation.remote
    source_expected_commit = source_attestation.expected_commit
    source_license = source_root / "LICENSE"
    if not source_license.is_file() or source_license.is_symlink():
        raise ValueError("source Harvey license is missing or unsafe")
    source_license_hash = source_attestation.license_sha256

    tasks_root = (source_root / "tasks").resolve()
    task_dir = source_root / "tasks" / Path(*task_relative.parts)
    ensure_safe_descendant(task_dir, tasks_root, "source task")
    task_path = task_dir / "task.json"
    if not task_path.is_file() or task_path.is_symlink():
        raise ValueError(f"no safe task.json at {task_path}")
    documents = source_documents(task_dir)
    documents_root = task_dir / "documents"
    source_task_hash = sha256_file(task_path)
    source_hashes = {path: sha256_file(path) for path in documents}

    if not entities_path.is_file() or entities_path.is_symlink():
        raise ValueError(f"entities config is missing or unsafe: {entities_path}")
    entities_path = entities_path.resolve()
    entities = validate_entities(json.loads(entities_path.read_text(encoding="utf-8")))
    serialized_entities = canonical_json(entities)
    entities_sha256 = sha256_bytes(serialized_entities)
    corpus = "\n".join(
        [task_path.read_text(encoding="utf-8")] + [extract_all_text(path) for path in documents]
    )
    pairs = build_replacement_map(entities, seed, corpus)

    output_relative = output_relative_for(task_name, seed)
    output_task_dir = output_root / Path(*output_relative.parts)
    ensure_safe_descendant(output_task_dir, output_root, "output task")
    if output_task_dir.exists() and not replace_generated:
        raise FileExistsError(f"output task already exists: {output_task_dir}")
    if output_task_dir.exists():
        validate_existing_target(output_task_dir, task_name, seed)

    output_mapping: dict[Path, PurePosixPath] = {}
    seen_outputs: set[str] = set()
    for source_path in documents:
        relative = source_path.relative_to(documents_root)
        output_path = mutated_relative_path(relative, pairs)
        if output_path.as_posix() in seen_outputs:
            raise ValueError(f"document rename collision at {output_path}")
        seen_outputs.add(output_path.as_posix())
        output_mapping[source_path] = output_path

    source_task = json.loads(task_path.read_text(encoding="utf-8"))
    generated_task, task_counts = build_expected_task(
        source_task,
        pairs,
        task_name,
        seed,
        source_commit,
        source_expected_commit,
        source_license_hash,
        entities_sha256,
    )

    output_task_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="harvey-mutation-", dir=output_task_dir.parent) as temporary:
        temporary_root = Path(temporary)
        staging = temporary_root / output_task_dir.name
        staging_documents = staging / "documents"
        staging_documents.mkdir(parents=True)
        file_rows: list[dict[str, Any]] = []

        for source_path, output_relative_path in output_mapping.items():
            destination = staging_documents / Path(*output_relative_path.parts)
            mutate_document(source_path, destination, pairs)
            extension = source_path.suffix.lower()
            file_rows.append({
                "source_path": source_path.relative_to(documents_root).as_posix(),
                "output_path": output_relative_path.as_posix(),
                "extension": extension,
                "source_sha256": source_hashes[source_path],
                "output_sha256": sha256_file(destination),
                "source_bytes": source_path.stat().st_size,
                "output_bytes": destination.stat().st_size,
                "package_topology_preserved": (
                    package_topology(source_path) == package_topology(destination)
                    if extension in OOXML_EXTENSIONS else None
                ),
            })

        (staging / "task.json").write_bytes(canonical_json(generated_task))
        (staging / "mutation-entities.json").write_bytes(serialized_entities)
        manifest = {
            "schema_version": 2,
            "tool": f"research/lab_mutate.py v{TOOL_VERSION}",
            "task": task_name,
            "output_task": output_relative.as_posix(),
            "seed": seed,
            "source_repo": SOURCE_REPO,
            "source_remote": source_remote,
            "source_commit": source_commit,
            "source_expected_commit": source_expected_commit,
            "source_git_clean": True,
            "source_license": "MIT",
            "source_license_path": "LICENSE",
            "source_license_sha256": source_license_hash,
            "source_task_json_sha256": source_task_hash,
            "output_task_json_sha256": sha256_file(staging / "task.json"),
            "entities_sha256": entities_sha256,
            "entity_count": len(entities),
            "substitutions": [{"old": old, "new": new} for old, new in pairs],
            "task_replacement_counts": dict(sorted(task_counts.items())),
            "files": sorted(file_rows, key=lambda row: row["source_path"]),
            "validation": {
                "source_entity_residuals": 0,
                "document_count_preserved": True,
                "source_immutable": True,
            },
        }
        (staging / "mutation-manifest.json").write_bytes(canonical_json(manifest))

        if sha256_file(task_path) != source_task_hash:
            raise RuntimeError("source task.json changed during generation")
        if sha256_file(source_license) != source_license_hash:
            raise RuntimeError("source Harvey license changed during generation")
        for source_path, source_hash in source_hashes.items():
            if sha256_file(source_path) != source_hash:
                raise RuntimeError(f"source changed during generation: {source_path}")
        if sha256_bytes(canonical_json(json.loads(entities_path.read_text(encoding="utf-8")))) != entities_sha256:
            raise RuntimeError("entities config changed during generation")
        if owns_attestation:
            verify_source_attestation(source_attestation)

        check_generated(
            staging,
            source_root,
            quiet=True,
            source_attestation=source_attestation,
        )
        if output_task_dir.exists():
            validate_existing_target(output_task_dir, task_name, seed)
            backup = temporary_root / "previous-generated-output"
            output_task_dir.replace(backup)
            try:
                staging.replace(output_task_dir)
            except Exception:
                if not output_task_dir.exists() and backup.exists():
                    backup.replace(output_task_dir)
                raise
        else:
            staging.replace(output_task_dir)

    return output_task_dir


def load_seed_plan(plan_path: Path) -> list[tuple[str, Path, list[int]]]:
    if not plan_path.is_file() or plan_path.is_symlink():
        raise ValueError(f"seed plan is missing or unsafe: {plan_path}")
    plan_path = plan_path.resolve()
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if not isinstance(plan, dict) or plan.get("schema_version") != 1:
        raise ValueError("seed plan schema_version must be 1")
    variants = plan.get("variants")
    if not isinstance(variants, list) or not variants:
        raise ValueError("seed plan variants must be a non-empty list")
    rows: list[tuple[str, Path, list[int]]] = []
    identities: set[tuple[str, int]] = set()
    for index, variant in enumerate(variants):
        if not isinstance(variant, dict):
            raise ValueError(f"seed plan variant {index} must be an object")
        task = variant.get("task")
        entities_value = variant.get("entities")
        seeds = variant.get("seeds")
        if not isinstance(task, str):
            raise ValueError(f"seed plan variant {index} requires task")
        task_relative = safe_relative(task, f"seed plan variant {index} task")
        if len(task_relative.parts) < 2:
            raise ValueError(f"seed plan variant {index} task needs practice-area/task-slug")
        if not isinstance(entities_value, str):
            raise ValueError(f"seed plan variant {index} requires entities path")
        entities_relative = safe_relative(
            entities_value, f"seed plan variant {index} entities"
        )
        entities_path = plan_path.parent / Path(*entities_relative.parts)
        ensure_safe_descendant(entities_path, plan_path.parent, "seed-plan entities")
        if not entities_path.is_file() or entities_path.is_symlink():
            raise ValueError(f"seed-plan entities config is missing or unsafe: {entities_path}")
        if (
            not isinstance(seeds, list)
            or not seeds
            or any(isinstance(seed, bool) or not isinstance(seed, int) or seed < 0 for seed in seeds)
            or len(seeds) != len(set(seeds))
        ):
            raise ValueError(f"seed plan variant {index} seeds must be unique non-negative integers")
        for seed in seeds:
            identity = (task, seed)
            if identity in identities:
                raise ValueError(f"duplicate seed-plan identity: {task} seed {seed}")
            identities.add(identity)
        rows.append((task, entities_path, seeds))
    return rows


def generate_seed_plan(
    plan_path: Path,
    source_root: Path,
    output_root: Path,
    replace_generated: bool = False,
    source_attestation: SourceAttestation | None = None,
) -> list[Path]:
    source_attestation, owns_attestation = use_source_attestation(
        source_root, source_attestation
    )
    outputs = []
    for task, entities_path, seeds in load_seed_plan(plan_path):
        for seed in seeds:
            outputs.append(
                generate(
                    task,
                    entities_path,
                    seed,
                    source_root,
                    output_root,
                    replace_generated,
                    source_attestation=source_attestation,
                )
            )
    if owns_attestation:
        verify_source_attestation(source_attestation)
    return outputs


def check_seed_plan(
    plan_path: Path = DEFAULT_PLAN,
    source_root: Path = LAB,
    output_root: Path = DEFAULT_OUTPUT,
    source_attestation: SourceAttestation | None = None,
) -> None:
    source_attestation, owns_attestation = use_source_attestation(
        source_root, source_attestation
    )
    plan_rows = load_seed_plan(plan_path)
    output_root = output_root.resolve()
    expected: dict[Path, Path] = {}
    for task, entities_path, seeds in plan_rows:
        serialized_entities = canonical_json(
            validate_entities(json.loads(entities_path.read_text(encoding="utf-8")))
        )
        for seed in seeds:
            output_relative = output_relative_for(task, seed)
            task_dir = output_root / Path(*output_relative.parts)
            ensure_safe_descendant(task_dir, output_root, "planned output task")
            if task_dir in expected:
                raise ValueError(f"seed plan targets output twice: {task_dir}")
            expected[task_dir] = entities_path
            copied_entities = task_dir / "mutation-entities.json"
            if not copied_entities.is_file() or copied_entities.read_bytes() != serialized_entities:
                raise ValueError(f"generated entity config differs from plan: {task_dir}")

    manifest_dirs = {
        path.parent.resolve()
        for path in output_root.rglob("mutation-manifest.json")
    } if output_root.is_dir() else set()
    if manifest_dirs != set(expected):
        missing = sorted(str(path) for path in set(expected) - manifest_dirs)
        orphaned = sorted(str(path) for path in manifest_dirs - set(expected))
        raise ValueError(f"seed-plan output mismatch; missing={missing}, orphaned={orphaned}")
    for task_dir in sorted(expected, key=str):
        check_generated(
            task_dir,
            source_root,
            quiet=True,
            source_attestation=source_attestation,
        )
    loose_files = [
        path for path in output_root.rglob("*")
        if path.is_file() and not any(path.is_relative_to(task_dir) for task_dir in expected)
    ] if output_root.is_dir() else []
    if loose_files:
        raise ValueError(
            "seed output root contains files outside planned tasks: "
            + ", ".join(str(path.relative_to(output_root)) for path in loose_files)
        )
    if owns_attestation:
        verify_source_attestation(source_attestation)
    print(
        f"seeded Harvey mutation plan valid: {len(expected)} variants, "
        f"source {source_attestation.commit[:12]}"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--task", help="area/slug relative to LAB tasks/")
    mode.add_argument("--check", type=Path, help="validate one generated seed directory")
    mode.add_argument("--plan", type=Path, help="generate every variant in a seed plan")
    mode.add_argument("--check-plan", type=Path, help="validate every output in a seed plan")
    parser.add_argument("--entities", type=Path)
    parser.add_argument("--seed", type=int)
    parser.add_argument("--source", type=Path, default=LAB)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--replace-generated",
        action="store_true",
        help="replace only an existing generated directory for this exact task and seed",
    )
    args = parser.parse_args()

    with mutation_output_lock(args.out):
        if args.check:
            if args.entities is not None or args.seed is not None or args.replace_generated:
                parser.error("--check cannot be combined with generation-only options")
            check_generated(args.check, args.source)
            return 0
        if args.check_plan:
            if args.entities is not None or args.seed is not None or args.replace_generated:
                parser.error("--check-plan cannot be combined with generation-only options")
            check_seed_plan(args.check_plan, args.source, args.out)
            return 0
        if args.plan:
            if args.entities is not None or args.seed is not None:
                parser.error("--plan reads entities and seeds from the plan file")
            source_attestation = attest_source(args.source)
            outputs = generate_seed_plan(
                args.plan,
                args.source,
                args.out,
                args.replace_generated,
                source_attestation=source_attestation,
            )
            for output in outputs:
                print(f"generated Harvey task: {output}")
            check_seed_plan(
                args.plan,
                args.source,
                args.out,
                source_attestation=source_attestation,
            )
            verify_source_attestation(source_attestation)
            return 0
        if args.entities is None or args.seed is None:
            parser.error("--task requires --entities and --seed")
        source_attestation = attest_source(args.source)
        output = generate(
            args.task,
            args.entities,
            args.seed,
            args.source,
            args.out,
            args.replace_generated,
            source_attestation=source_attestation,
        )
        print(f"generated Harvey task: {output}")
        check_generated(
            output,
            args.source,
            source_attestation=source_attestation,
        )
        verify_source_attestation(source_attestation)
        return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (
        FileExistsError,
        KeyError,
        OSError,
        RuntimeError,
        subprocess.CalledProcessError,
        TypeError,
        UnicodeDecodeError,
        ValueError,
        zipfile.BadZipFile,
    ) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
