#!/usr/bin/env python3
"""Create deterministic, provenance-rich Harvey LAB task derivatives.

The tool never edits the Harvey mirror.  It copies one task into a separate
augmentation tree, applies fail-closed text substitutions to text/OOXML parts,
updates task metadata and rubric criteria from a declarative recipe, then
records source/output hashes and package-topology checks in a manifest.

Text substitutions in OOXML are performed on raw XML bytes.  The XML is not
parsed and reserialized, so run, style, relationship, drawing, worksheet, and
section structures remain byte-for-byte untouched outside matched text.

Usage:
    python3 tools/mutate_harvey_task.py --recipe path/to/recipe.json
    python3 tools/mutate_harvey_task.py --check path/to/generated/task
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from email import policy
from email.parser import BytesParser
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree
from xml.sax.saxutils import escape as xml_escape


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "research" / "repos" / "harveyai@harvey-labs"
DEFAULT_OUTPUT = ROOT / "research" / "harvey-augmentation" / "generated"
DEFAULT_RECIPES = ROOT / "research" / "harvey-augmentation" / "recipes"
DEFAULT_PIN_FILE = ROOT / "research" / "repos-commits.json"
PIN_KEY = "harveyai@harvey-labs"
SOURCE_REPO = "harveyai/harvey-labs"
OOXML_EXTENSIONS = {".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"}
TEXT_EXTENSIONS = {".eml", ".txt", ".md", ".json", ".csv", ".xml", ".html", ".htm"}
XML_PART_EXTENSIONS = {".xml", ".rels"}


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_xml_part(name: str) -> bool:
    return name.lower().endswith((".xml", ".rels"))


def canonical_json(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def safe_relative(value: str, label: str) -> PurePosixPath:
    path = PurePosixPath(value)
    if (
        path.is_absolute()
        or not path.parts
        or "\\" in value
        or any(ord(character) < 32 for character in value)
        or any(part in {"", ".", ".."} for part in path.parts)
    ):
        raise ValueError(f"unsafe {label}: {value!r}")
    return path


def ensure_safe_descendant(path: Path, root: Path, label: str) -> Path:
    root = root.resolve()
    try:
        relative = path.relative_to(root)
        resolved = path.resolve(strict=False)
        resolved.relative_to(root)
    except ValueError as error:
        raise ValueError(f"{label} escapes its root: {path}") from error
    cursor = root
    for part in relative.parts:
        cursor /= part
        if cursor.is_symlink():
            raise ValueError(f"{label} traverses a symlink: {cursor}")
    return resolved


def git_output(source: Path, *arguments: str) -> str:
    return subprocess.run(
        ["git", "-C", str(source), *arguments],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()


def git_commit(source: Path) -> str:
    return git_output(source, "rev-parse", "HEAD")


def git_clean(source: Path) -> bool:
    return not bool(git_output(source, "status", "--porcelain", "--untracked-files=all"))


def git_remote(source: Path) -> str:
    return git_output(source, "remote", "get-url", "origin")


def normalized_remote(remote: str) -> str:
    value = remote.strip().rstrip("/")
    if value.endswith(".git"):
        value = value[:-4]
    if value.startswith("git@github.com:"):
        value = "https://github.com/" + value.removeprefix("git@github.com:")
    return value.lower()


def validate_source_identity(source: Path, commit: str, remote: str) -> str | None:
    if source.resolve() != DEFAULT_SOURCE.resolve():
        return None
    if normalized_remote(remote) != "https://github.com/harveyai/harvey-labs":
        raise ValueError(f"default Harvey mirror has unexpected origin: {remote}")
    pins = json.loads(DEFAULT_PIN_FILE.read_text(encoding="utf-8"))
    expected = pins.get(PIN_KEY)
    if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-fA-F]{7,40}", expected):
        raise ValueError(f"missing or invalid Harvey source pin in {DEFAULT_PIN_FILE}")
    if not commit.startswith(expected):
        raise ValueError(f"source commit mismatch: expected {expected}, got {commit}")
    return expected


def validate_recipe(recipe: dict[str, Any]) -> None:
    if recipe.get("schema_version") != 1:
        raise ValueError("recipe schema_version must be 1")
    for field in ("variant_id", "source_task", "output_task"):
        if not isinstance(recipe.get(field), str) or not recipe[field].strip():
            raise ValueError(f"recipe requires non-empty {field}")
        if any(ord(character) < 32 for character in recipe[field]):
            raise ValueError(f"recipe {field} contains control characters")
    source_task = safe_relative(recipe["source_task"], "source_task")
    output_task = safe_relative(recipe["output_task"], "output_task")
    if len(source_task.parts) < 2 or len(output_task.parts) < 2:
        raise ValueError("source_task and output_task require at least practice-area/task-slug")

    replacements = recipe.get("replacements") or []
    if not isinstance(replacements, list):
        raise ValueError("recipe replacements must be a list")
    seen = set()
    for index, replacement in enumerate(replacements):
        if not isinstance(replacement, dict):
            raise ValueError(f"replacement {index} must be an object")
        old = replacement.get("old")
        new = replacement.get("new")
        if not isinstance(old, str) or not old:
            raise ValueError(f"replacement {index} requires non-empty old text")
        if not isinstance(new, str):
            raise ValueError(f"replacement {index} requires string new text")
        if old in seen:
            raise ValueError(f"duplicate replacement source: {old!r}")
        seen.add(old)
        if old == new:
            raise ValueError(f"replacement {index} does not change text")
        if recipe.get("layout_preserving", True):
            if len(old.encode("utf-8")) != len(new.encode("utf-8")):
                raise ValueError(
                    f"layout-preserving recipe requires equal UTF-8 byte length: {old!r} -> {new!r}"
                )
            if len(xml_escape(old).encode("utf-8")) != len(xml_escape(new).encode("utf-8")):
                raise ValueError(
                    f"layout-preserving recipe requires equal XML-escaped byte length: "
                    f"{old!r} -> {new!r}"
                )
        extensions = replacement.get("extensions")
        if extensions is not None:
            if (
                not isinstance(extensions, list)
                or not extensions
                or any(not isinstance(value, str) or not value.startswith(".") for value in extensions)
            ):
                raise ValueError(f"replacement {index} extensions must be non-empty dotted strings")
        for count_field in ("min_document_occurrences", "min_task_occurrences"):
            count = replacement.get(count_field, 0)
            if isinstance(count, bool) or not isinstance(count, int) or count < 0:
                raise ValueError(f"replacement {index} {count_field} must be a non-negative integer")

    for old in seen:
        for replacement in replacements:
            if old in replacement["new"]:
                raise ValueError(
                    f"replacement output {replacement['new']!r} contains source term {old!r}"
                )
    for mapping_field in ("task_overrides", "criterion_overrides"):
        mapping = recipe.get(mapping_field, {})
        if not isinstance(mapping, dict):
            raise ValueError(f"recipe {mapping_field} must be an object")
        if any(not isinstance(key, str) or not key for key in mapping):
            raise ValueError(f"recipe {mapping_field} keys must be non-empty strings")
    if any(
        not isinstance(value, dict)
        for value in (recipe.get("criterion_overrides") or {}).values()
    ):
        raise ValueError("recipe criterion_overrides values must be objects")
    axes = recipe.get("mutation_axes", [])
    if not isinstance(axes, list) or any(not isinstance(axis, str) or not axis for axis in axes):
        raise ValueError("recipe mutation_axes must be a list of non-empty strings")
    expected_commit = recipe.get("source_commit")
    if expected_commit is not None and (
        not isinstance(expected_commit, str)
        or not re.fullmatch(r"[0-9a-fA-F]{7,40}", expected_commit)
    ):
        raise ValueError("recipe source_commit must be a 7-40 character hexadecimal prefix")


def deep_replace(value: Any, replacements: list[dict[str, Any]], counts: Counter[str]) -> Any:
    if isinstance(value, str):
        for replacement in replacements:
            old, new = replacement["old"], replacement["new"]
            occurrences = value.count(old)
            if occurrences:
                value = value.replace(old, new)
                counts[old] += occurrences
        return value
    if isinstance(value, list):
        return [deep_replace(item, replacements, counts) for item in value]
    if isinstance(value, dict):
        return {key: deep_replace(item, replacements, counts) for key, item in value.items()}
    return value


def replacement_applies(replacement: dict[str, Any], extension: str) -> bool:
    extensions = replacement.get("extensions")
    return not extensions or extension in {str(value).lower() for value in extensions}


def package_topology(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        rows = [
            {
                "name": info.filename,
                "compression": info.compress_type,
                "date_time": info.date_time,
                "flag_bits": info.flag_bits,
                "create_system": info.create_system,
                "create_version": info.create_version,
                "extract_version": info.extract_version,
                "external_attr": info.external_attr,
                "internal_attr": info.internal_attr,
                "extra": info.extra.hex(),
                "comment": info.comment.hex(),
                "directory": info.is_dir(),
            }
            for info in archive.infolist()
        ]
        archive_comment = archive.comment.hex()
    return {
        "parts": len(rows),
        "archive_comment": archive_comment,
        "sha256": sha256_bytes(canonical_json(rows)),
    }


def validate_ooxml(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"ZIP CRC failure in {bad}")
        for name in archive.namelist():
            if is_xml_part(name):
                ElementTree.fromstring(archive.read(name))
    extension = path.suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        import openpyxl

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_vba=extension == ".xlsm",
        )
        try:
            if not workbook.sheetnames:
                raise ValueError(f"workbook has no worksheets: {path.name}")
        finally:
            workbook.close()
    elif extension in {".docx", ".docm"}:
        from docx import Document

        Document(path)
    elif extension in {".pptx", ".pptm"}:
        from pptx import Presentation

        Presentation(path)


def ooxml_semantic_text(path: Path) -> str:
    """Extract attributes, text nodes, and run-joined paragraph/string text."""
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not is_xml_part(name):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for element in root.iter():
                chunks.extend(str(value) for value in element.attrib.values())
                if element.text:
                    chunks.append(element.text)
                local_name = element.tag.rsplit("}", 1)[-1]
                if local_name in {"p", "si", "is", "txBody"}:
                    joined = "".join(element.itertext())
                    if joined:
                        chunks.append(joined)
    return "\n".join(chunks)


def document_semantic_text(path: Path) -> str | None:
    extension = path.suffix.lower()
    if extension in OOXML_EXTENSIONS:
        return ooxml_semantic_text(path)
    if extension == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        chunks = [f"{name}: {value}" for name, value in message.items()]
        for part in message.walk():
            if not part.is_multipart() and part.get_content_maintype() == "text":
                content = part.get_content()
                if isinstance(content, str):
                    chunks.append(content)
        return "\n".join(chunks)
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8")
    return None


def assert_no_residuals(path: Path, replacements: list[dict[str, Any]]) -> None:
    applicable = [item for item in replacements if replacement_applies(item, path.suffix.lower())]
    if not applicable:
        return
    text = document_semantic_text(path)
    if text is None:
        return
    residuals = sorted({item["old"] for item in applicable if item["old"] in text})
    if residuals:
        raise ValueError(f"source replacement text remains in {path.name}: {residuals}")


def mutate_ooxml(
    source: Path,
    destination: Path,
    replacements: list[dict[str, Any]],
) -> tuple[Counter[str], list[str]]:
    counts: Counter[str] = Counter()
    changed_parts: list[str] = []
    with zipfile.ZipFile(source) as input_archive, zipfile.ZipFile(destination, "w") as output_archive:
        output_archive.comment = input_archive.comment
        for info in input_archive.infolist():
            payload = input_archive.read(info.filename)
            original = payload
            if is_xml_part(info.filename):
                for replacement in replacements:
                    old = xml_escape(replacement["old"]).encode("utf-8")
                    new = xml_escape(replacement["new"]).encode("utf-8")
                    occurrences = payload.count(old)
                    if occurrences:
                        payload = payload.replace(old, new)
                        counts[replacement["old"]] += occurrences
            if payload != original:
                changed_parts.append(info.filename)
            output_archive.writestr(info, payload)
    if package_topology(source) != package_topology(destination):
        raise ValueError(f"OOXML package topology changed: {source.name}")
    validate_ooxml(destination)
    return counts, changed_parts


def mutate_text(
    source: Path,
    destination: Path,
    replacements: list[dict[str, Any]],
) -> tuple[Counter[str], list[str]]:
    payload = source.read_bytes()
    counts: Counter[str] = Counter()
    for replacement in replacements:
        old = replacement["old"].encode("utf-8")
        new = replacement["new"].encode("utf-8")
        occurrences = payload.count(old)
        if occurrences:
            payload = payload.replace(old, new)
            counts[replacement["old"]] += occurrences
    destination.write_bytes(payload)
    if destination.suffix.lower() == ".json":
        json.loads(destination.read_text(encoding="utf-8"))
    elif destination.suffix.lower() == ".eml":
        with destination.open("rb") as handle:
            BytesParser(policy=policy.default).parse(handle)
    return counts, [destination.name] if counts else []


def copy_or_mutate(
    source: Path,
    destination: Path,
    replacements: list[dict[str, Any]],
) -> tuple[Counter[str], list[str]]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    extension = source.suffix.lower()
    applicable = [item for item in replacements if replacement_applies(item, extension)]
    if not applicable:
        shutil.copy2(source, destination)
        return Counter(), []
    if extension in OOXML_EXTENSIONS:
        result = mutate_ooxml(source, destination, applicable)
        assert_no_residuals(destination, applicable)
        return result
    if extension in TEXT_EXTENSIONS:
        result = mutate_text(source, destination, applicable)
        assert_no_residuals(destination, applicable)
        return result

    raise ValueError(
        f"refusing unsupported format with applicable replacements: {source.name}"
    )


def verify_document_derivative(
    source: Path,
    output: Path,
    replacements: list[dict[str, Any]],
) -> tuple[Counter[str], list[str]]:
    """Re-derive expected content and reject any unrecorded document change."""
    extension = source.suffix.lower()
    applicable = [item for item in replacements if replacement_applies(item, extension)]
    counts: Counter[str] = Counter()
    changed_parts: list[str] = []

    if extension in OOXML_EXTENSIONS:
        with zipfile.ZipFile(source) as source_archive, zipfile.ZipFile(output) as output_archive:
            if source_archive.namelist() != output_archive.namelist():
                raise ValueError(f"OOXML part order differs from source: {source.name}")
            if source_archive.comment != output_archive.comment:
                raise ValueError(f"OOXML archive comment differs from source: {source.name}")
            for name in source_archive.namelist():
                expected = source_archive.read(name)
                original = expected
                if is_xml_part(name):
                    for replacement in applicable:
                        old = xml_escape(replacement["old"]).encode("utf-8")
                        new = xml_escape(replacement["new"]).encode("utf-8")
                        occurrences = expected.count(old)
                        if occurrences:
                            expected = expected.replace(old, new)
                            counts[replacement["old"]] += occurrences
                if expected != original:
                    changed_parts.append(name)
                if output_archive.read(name) != expected:
                    raise ValueError(
                        f"OOXML content has a non-recipe change: {source.name}:{name}"
                    )
        assert_no_residuals(output, applicable)
        return counts, changed_parts

    if applicable and extension not in TEXT_EXTENSIONS:
        raise ValueError(f"unsupported format has applicable replacements: {source.name}")

    expected = source.read_bytes()
    original = expected
    if extension in TEXT_EXTENSIONS:
        for replacement in applicable:
            old = replacement["old"].encode("utf-8")
            new = replacement["new"].encode("utf-8")
            occurrences = expected.count(old)
            if occurrences:
                expected = expected.replace(old, new)
                counts[replacement["old"]] += occurrences
    if output.read_bytes() != expected:
        raise ValueError(f"document content has a non-recipe change: {source.name}")
    if expected != original:
        changed_parts.append(output.name)
    assert_no_residuals(output, applicable)
    return counts, changed_parts


def validate_task(task: dict[str, Any]) -> None:
    if not isinstance(task, dict):
        raise ValueError("generated task must be an object")
    if not isinstance(task.get("title"), str) or not task["title"].strip():
        raise ValueError("generated task title is empty")
    if not isinstance(task.get("instructions"), str) or not task["instructions"].strip():
        raise ValueError("generated task instructions are empty")
    deliverables = task.get("deliverables")
    if not isinstance(deliverables, dict) or not deliverables:
        raise ValueError("generated task requires deliverables mapping")
    deliverable_names: set[str] = set()
    for key, value in deliverables.items():
        if not isinstance(key, str) or not key.strip():
            raise ValueError("generated deliverable keys must be non-empty strings")
        safe_relative(key, "deliverable key")
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"generated deliverable {key!r} must map to a non-empty filename")
        safe_relative(value, f"deliverable {key!r}")
        deliverable_names.update({key, value})
    criteria = task.get("criteria")
    if not isinstance(criteria, list) or not criteria:
        raise ValueError("generated task requires criteria")
    if any(not isinstance(criterion, dict) for criterion in criteria):
        raise ValueError("generated criteria must be objects")
    identifiers = [criterion.get("id") for criterion in criteria]
    if (
        any(not isinstance(identifier, str) or not identifier.strip() for identifier in identifiers)
        or len(identifiers) != len(set(identifiers))
    ):
        raise ValueError("generated criterion IDs must be non-empty and unique")
    for criterion in criteria:
        if not isinstance(criterion.get("title"), str) or not criterion["title"].strip():
            raise ValueError(f"criterion {criterion.get('id')} requires a non-empty title")
        criterion_deliverables = criterion.get("deliverables")
        if not isinstance(criterion_deliverables, list) or not criterion_deliverables:
            raise ValueError(f"criterion {criterion.get('id')} requires a deliverables list")
        if any(
            not isinstance(value, str) or value not in deliverable_names
            for value in criterion_deliverables
        ):
            raise ValueError(
                f"criterion {criterion.get('id')} references an unknown deliverable"
            )
        match_criteria = criterion.get("match_criteria")
        if not isinstance(match_criteria, str):
            raise ValueError(f"criterion {criterion.get('id')} match_criteria must be text")
        if not match_criteria.startswith("PASS if") or "FAIL if" not in match_criteria:
            raise ValueError(f"criterion {criterion.get('id')} requires fail-closed PASS if text")


def apply_task_recipe(
    source_task: dict[str, Any],
    recipe: dict[str, Any],
    source_commit: str,
    recipe_sha256: str,
    source_expected_commit: str | None = None,
) -> tuple[dict[str, Any], Counter[str]]:
    counts: Counter[str] = Counter()
    task = deep_replace(source_task, recipe.get("replacements") or [], counts)
    for key, value in (recipe.get("task_overrides") or {}).items():
        task[key] = value

    overrides = recipe.get("criterion_overrides") or {}
    found = set()
    for criterion in task.get("criteria") or []:
        identifier = criterion.get("id")
        if identifier in overrides:
            criterion.update(overrides[identifier])
            found.add(identifier)
    missing = sorted(set(overrides) - found)
    if missing:
        raise ValueError(f"criterion overrides reference missing IDs: {missing}")

    task["augmentation"] = {
        "variant_id": recipe["variant_id"],
        "source_repo": SOURCE_REPO,
        "source_license": "MIT",
        "source_commit": source_commit,
        "source_expected_commit": source_expected_commit,
        "source_task": recipe["source_task"],
        "recipe_sha256": recipe_sha256,
        "mutation_axes": recipe.get("mutation_axes") or [],
        "originals_immutable": True,
    }
    validate_task(task)
    return task, counts


def collect_document_paths(documents_root: Path) -> list[Path]:
    if not documents_root.is_dir() or documents_root.is_symlink():
        raise ValueError(f"documents directory is missing or unsafe: {documents_root}")
    if any(path.is_symlink() for path in documents_root.rglob("*")):
        raise ValueError(f"documents tree contains a symlink: {documents_root}")
    paths = sorted(
        (path for path in documents_root.rglob("*") if path.is_file()),
        key=lambda path: path.relative_to(documents_root).as_posix(),
    )
    if not paths:
        raise ValueError(f"documents directory is empty: {documents_root}")
    for path in paths:
        if path.is_symlink():
            raise ValueError(f"source document is a symlink: {path}")
    return paths


def validate_existing_target(target: Path, recipe: dict[str, Any]) -> None:
    if target.is_symlink():
        raise ValueError(f"refusing to replace symlink output: {target}")
    manifest_path = target / "mutation-manifest.json"
    if not manifest_path.is_file() or manifest_path.is_symlink():
        raise ValueError(f"refusing to replace unrecognized output directory: {target}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for field in ("variant_id", "source_task", "output_task"):
        if manifest.get(field) != recipe[field]:
            raise ValueError(f"existing output manifest disagrees on {field}")


def generate(
    recipe_path: Path,
    source_root: Path,
    output_root: Path,
    replace_generated: bool = False,
) -> Path:
    if not recipe_path.is_file() or recipe_path.is_symlink():
        raise ValueError(f"augmentation recipe is missing or unsafe: {recipe_path}")
    recipe_path = recipe_path.resolve()
    source_root = source_root.resolve()
    output_root = output_root.resolve()
    recipe = json.loads(recipe_path.read_text(encoding="utf-8"))
    validate_recipe(recipe)
    serialized_recipe = canonical_json(recipe)
    recipe_sha256 = sha256_bytes(serialized_recipe)
    if not git_clean(source_root):
        raise ValueError("source Harvey worktree is dirty")
    source_commit = git_commit(source_root)
    source_remote = git_remote(source_root)
    source_expected_commit = validate_source_identity(
        source_root, source_commit, source_remote
    )
    source_license = source_root / "LICENSE"
    if not source_license.is_file() or source_license.is_symlink():
        raise ValueError("source Harvey license is missing or unsafe")
    source_license_hash = sha256_file(source_license)
    expected_commit = recipe.get("source_commit")
    if expected_commit and not source_commit.startswith(expected_commit):
        raise ValueError(f"source commit mismatch: expected {expected_commit}, got {source_commit}")
    if source_expected_commit is not None and expected_commit != source_expected_commit:
        raise ValueError(
            "recipe source_commit must exactly match the checked-in Harvey source pin"
        )

    source_task_relative = safe_relative(recipe["source_task"], "source_task")
    output_task_relative = safe_relative(recipe["output_task"], "output_task")
    source_task_dir = source_root / "tasks" / Path(*source_task_relative.parts)
    ensure_safe_descendant(source_task_dir, (source_root / "tasks").resolve(), "source task")
    source_task_json = source_task_dir / "task.json"
    source_documents = source_task_dir / "documents"
    if (
        not source_task_json.is_file()
        or source_task_json.is_symlink()
        or not source_documents.is_dir()
        or source_documents.is_symlink()
    ):
        raise ValueError("pilot generator currently requires a task-local documents directory")

    output_task_dir = output_root / "tasks" / Path(*output_task_relative.parts)
    ensure_safe_descendant(output_task_dir, (output_root / "tasks").resolve(), "output task")
    if output_task_dir.exists() and not replace_generated:
        raise FileExistsError(f"output task already exists: {output_task_dir}")
    if output_task_dir.exists():
        validate_existing_target(output_task_dir, recipe)

    source_paths = collect_document_paths(source_documents)
    source_hashes = {path: sha256_file(path) for path in source_paths}
    source_task_hash = sha256_file(source_task_json)
    source_task = json.loads(source_task_json.read_text(encoding="utf-8"))
    generated_task, task_counts = apply_task_recipe(
        source_task,
        recipe,
        source_commit,
        recipe_sha256,
        source_expected_commit,
    )

    output_task_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="harvey-augmentation-", dir=output_task_dir.parent) as temporary:
        staging = Path(temporary) / output_task_dir.name
        staging_documents = staging / "documents"
        staging_documents.mkdir(parents=True)
        aggregate_counts: Counter[str] = Counter()
        file_rows = []

        for source_path in source_paths:
            relative = source_path.relative_to(source_documents)
            path_residuals = [
                item["old"]
                for item in recipe.get("replacements") or []
                if replacement_applies(item, source_path.suffix.lower())
                and item["old"] in relative.as_posix()
            ]
            if path_residuals:
                raise ValueError(
                    f"strict generator does not rename document paths; {relative} contains "
                    f"source text {sorted(set(path_residuals))}"
                )
            destination = staging_documents / relative
            counts, changed_parts = copy_or_mutate(
                source_path, destination, recipe.get("replacements") or []
            )
            aggregate_counts.update(counts)
            file_rows.append({
                "path": relative.as_posix(),
                "extension": source_path.suffix.lower(),
                "source_sha256": source_hashes[source_path],
                "output_sha256": sha256_file(destination),
                "source_bytes": source_path.stat().st_size,
                "output_bytes": destination.stat().st_size,
                "replacement_counts": dict(sorted(counts.items())),
                "changed_package_parts": sorted(changed_parts),
                "package_topology_preserved": (
                    package_topology(source_path) == package_topology(destination)
                    if source_path.suffix.lower() in OOXML_EXTENSIONS else None
                ),
            })

        for replacement in recipe.get("replacements") or []:
            old = replacement["old"]
            if aggregate_counts[old] < int(replacement.get("min_document_occurrences", 0)):
                raise ValueError(
                    f"document replacement count for {old!r} is {aggregate_counts[old]}, "
                    f"expected at least {replacement.get('min_document_occurrences', 0)}"
                )
            if task_counts[old] < int(replacement.get("min_task_occurrences", 0)):
                raise ValueError(
                    f"task replacement count for {old!r} is {task_counts[old]}, "
                    f"expected at least {replacement.get('min_task_occurrences', 0)}"
                )

        for source_path, digest in source_hashes.items():
            if sha256_file(source_path) != digest:
                raise RuntimeError(f"source changed during generation: {source_path}")
        if sha256_file(source_task_json) != source_task_hash:
            raise RuntimeError("source task.json changed during generation")
        if sha256_file(source_license) != source_license_hash:
            raise RuntimeError("source license changed during generation")
        if collect_document_paths(source_documents) != source_paths:
            raise RuntimeError("source document path set changed during generation")
        if git_commit(source_root) != source_commit or not git_clean(source_root):
            raise RuntimeError("source Harvey worktree changed during generation")

        (staging / "task.json").write_bytes(canonical_json(generated_task))
        serialized_task = (staging / "task.json").read_text(encoding="utf-8")
        task_residuals = sorted({
            replacement["old"]
            for replacement in recipe.get("replacements") or []
            if replacement["old"] in serialized_task
        })
        if task_residuals:
            raise ValueError(f"source replacement text remains in generated task: {task_residuals}")
        (staging / "augmentation-recipe.json").write_bytes(serialized_recipe)
        manifest = {
            "schema_version": 2,
            "variant_id": recipe["variant_id"],
            "source_repo": SOURCE_REPO,
            "source_remote": source_remote,
            "source_commit": source_commit,
            "source_expected_commit": source_expected_commit,
            "source_license": "MIT",
            "source_license_path": "LICENSE",
            "source_license_sha256": source_license_hash,
            "source_task": recipe["source_task"],
            "output_task": recipe["output_task"],
            "recipe_sha256": recipe_sha256,
            "source_task_json_sha256": source_task_hash,
            "output_task_json_sha256": sha256_file(staging / "task.json"),
            "document_replacement_counts": dict(sorted(aggregate_counts.items())),
            "task_replacement_counts": dict(sorted(task_counts.items())),
            "files": file_rows,
        }
        (staging / "mutation-manifest.json").write_bytes(canonical_json(manifest))
        check_generated(staging, source_root, quiet=True)
        if output_task_dir.exists():
            validate_existing_target(output_task_dir, recipe)
            backup = Path(temporary) / "previous-generated-output"
            output_task_dir.replace(backup)
            try:
                staging.replace(output_task_dir)
            except Exception:
                if not output_task_dir.exists() and backup.exists():
                    backup.replace(output_task_dir)
                raise
        else:
            staging.replace(output_task_dir)

    return output_task_dir


def check_generated(
    task_dir: Path,
    source_root: Path = DEFAULT_SOURCE,
    quiet: bool = False,
) -> None:
    if task_dir.is_symlink():
        raise ValueError(f"generated task directory is a symlink: {task_dir}")
    task_dir = task_dir.resolve()
    source_root = source_root.resolve()
    manifest_path = task_dir / "mutation-manifest.json"
    task_path = task_dir / "task.json"
    recipe_path = task_dir / "augmentation-recipe.json"
    if any(
        not path.is_file() or path.is_symlink()
        for path in (manifest_path, task_path, recipe_path)
    ):
        raise ValueError(f"not a generated Harvey task: {task_dir}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    task = json.loads(task_path.read_text(encoding="utf-8"))
    recipe = json.loads(recipe_path.read_text(encoding="utf-8"))
    if not isinstance(manifest, dict) or manifest.get("schema_version") != 2:
        raise ValueError("generated manifest schema_version must be 2")
    for field in (
        "variant_id",
        "source_repo",
        "source_remote",
        "source_commit",
        "source_expected_commit",
        "source_task",
        "output_task",
        "recipe_sha256",
        "source_task_json_sha256",
        "output_task_json_sha256",
        "document_replacement_counts",
        "task_replacement_counts",
        "files",
    ):
        if field not in manifest:
            raise ValueError(f"generated manifest is missing {field}")
    if manifest["source_repo"] != SOURCE_REPO:
        raise ValueError("generated manifest source_repo is invalid")
    validate_recipe(recipe)
    validate_task(task)

    if sha256_file(recipe_path) != manifest["recipe_sha256"]:
        raise ValueError("generated recipe differs from its manifest")
    if sha256_file(task_path) != manifest["output_task_json_sha256"]:
        raise ValueError("generated task.json differs from its manifest")

    for field in ("variant_id", "source_task", "output_task"):
        if recipe[field] != manifest[field]:
            raise ValueError(f"recipe and manifest disagree on {field}")
    augmentation = task.get("augmentation") or {}
    for field in (
        "variant_id",
        "source_task",
        "source_commit",
        "source_expected_commit",
        "recipe_sha256",
    ):
        if augmentation.get(field) != manifest[field]:
            raise ValueError(f"task augmentation and manifest disagree on {field}")
    if augmentation.get("source_license") != "MIT":
        raise ValueError("task augmentation is missing MIT source-license provenance")
    if augmentation.get("source_repo") != SOURCE_REPO or augmentation.get("originals_immutable") is not True:
        raise ValueError("task augmentation has invalid source provenance")

    if not git_clean(source_root):
        raise ValueError("source Harvey worktree is dirty")
    current_commit = git_commit(source_root)
    if current_commit != manifest["source_commit"]:
        raise ValueError(
            f"source mirror moved: expected {manifest['source_commit']}, got {current_commit}"
        )
    current_remote = git_remote(source_root)
    if current_remote != manifest["source_remote"]:
        raise ValueError("source mirror remote differs from manifest")
    expected_commit = validate_source_identity(source_root, current_commit, current_remote)
    if manifest["source_expected_commit"] != expected_commit:
        raise ValueError("source pin differs from generated manifest")
    recipe_expected_commit = recipe.get("source_commit")
    if expected_commit is not None and recipe_expected_commit != expected_commit:
        raise ValueError("generated recipe does not match the checked-in Harvey source pin")
    if manifest.get("source_license") != "MIT" or manifest.get("source_license_path") != "LICENSE":
        raise ValueError("generated manifest has invalid source license metadata")
    source_license = source_root / "LICENSE"
    if (
        not source_license.is_file()
        or source_license.is_symlink()
        or sha256_file(source_license) != manifest.get("source_license_sha256")
    ):
        raise ValueError("source Harvey license differs from manifest")
    source_task_relative = safe_relative(manifest["source_task"], "source_task")
    source_task_dir = source_root / "tasks" / Path(*source_task_relative.parts)
    ensure_safe_descendant(source_task_dir, (source_root / "tasks").resolve(), "source task")
    source_task_json = source_task_dir / "task.json"
    source_documents = source_task_dir / "documents"
    if (
        not source_task_json.is_file()
        or source_task_json.is_symlink()
        or not source_documents.is_dir()
        or source_documents.is_symlink()
    ):
        raise ValueError("manifest source task or documents are missing")
    if sha256_file(source_task_json) != manifest["source_task_json_sha256"]:
        raise ValueError("source task.json differs from the generation-time hash")

    source_task = json.loads(source_task_json.read_text(encoding="utf-8"))
    expected_task, expected_task_counts = apply_task_recipe(
        source_task,
        recipe,
        manifest["source_commit"],
        manifest["recipe_sha256"],
        manifest["source_expected_commit"],
    )
    if canonical_json(expected_task) != task_path.read_bytes():
        raise ValueError("generated task.json contains a non-recipe change")
    task_residuals = sorted({
        replacement["old"]
        for replacement in recipe.get("replacements") or []
        if replacement["old"] in task_path.read_text(encoding="utf-8")
    })
    if task_residuals:
        raise ValueError(f"source replacement text remains in generated task: {task_residuals}")
    if dict(sorted(expected_task_counts.items())) != manifest["task_replacement_counts"]:
        raise ValueError("task replacement counts differ from the manifest")

    if not isinstance(manifest["files"], list) or any(
        not isinstance(row, dict) or not isinstance(row.get("path"), str)
        for row in manifest["files"]
    ):
        raise ValueError("manifest files must be a list of path rows")
    rows = {row["path"]: row for row in manifest["files"]}
    if len(rows) != len(manifest["files"]):
        raise ValueError("manifest contains duplicate document paths")
    current_sources = {
        path.relative_to(source_documents).as_posix(): path
        for path in collect_document_paths(source_documents)
    }
    if set(rows) != set(current_sources):
        raise ValueError("source document paths differ from generation manifest")
    if not (task_dir / "documents").is_dir() or (task_dir / "documents").is_symlink():
        raise ValueError("generated documents directory is missing or unsafe")
    actual = {
        path.relative_to(task_dir / "documents").as_posix(): path
        for path in (task_dir / "documents").rglob("*")
        if path.is_file()
    }
    if any(path.is_symlink() for path in task_dir.rglob("*")):
        raise ValueError("generated documents contain a symlink")
    if set(rows) != set(actual):
        raise ValueError("generated document paths differ from manifest")
    aggregate_counts: Counter[str] = Counter()
    for relative, path in sorted(actual.items()):
        row = rows[relative]
        path_residuals = [
            item["old"]
            for item in recipe.get("replacements") or []
            if replacement_applies(item, path.suffix.lower()) and item["old"] in relative
        ]
        if path_residuals:
            raise ValueError(
                f"generated document path contains source text: {relative}: "
                f"{sorted(set(path_residuals))}"
            )
        safe_path = safe_relative(relative, "manifest document path")
        source_path = source_documents / Path(*safe_path.parts)
        if not source_path.is_file():
            raise ValueError(f"source document is missing: {relative}")
        if sha256_file(source_path) != row["source_sha256"]:
            raise ValueError(f"source document differs from manifest: {relative}")
        if sha256_file(path) != row["output_sha256"]:
            raise ValueError(f"generated document differs from manifest: {relative}")
        if row.get("extension") != source_path.suffix.lower():
            raise ValueError(f"manifest extension is incorrect: {relative}")
        if row.get("source_bytes") != source_path.stat().st_size:
            raise ValueError(f"manifest source byte count is incorrect: {relative}")
        if row.get("output_bytes") != path.stat().st_size:
            raise ValueError(f"manifest output byte count is incorrect: {relative}")
        counts, changed_parts = verify_document_derivative(
            source_path, path, recipe.get("replacements") or []
        )
        aggregate_counts.update(counts)
        if dict(sorted(counts.items())) != row["replacement_counts"]:
            raise ValueError(f"document replacement counts differ from manifest: {relative}")
        if sorted(changed_parts) != row["changed_package_parts"]:
            raise ValueError(f"changed document parts differ from manifest: {relative}")
        if path.suffix.lower() in OOXML_EXTENSIONS:
            if row["package_topology_preserved"] is not True:
                raise ValueError(f"manifest does not confirm OOXML topology: {relative}")
            if package_topology(source_path) != package_topology(path):
                raise ValueError(f"OOXML package topology differs from source: {relative}")
            validate_ooxml(source_path)
            validate_ooxml(path)
        elif row.get("package_topology_preserved") is not None:
            raise ValueError(f"manifest has unexpected OOXML topology value: {relative}")
    if dict(sorted(aggregate_counts.items())) != manifest["document_replacement_counts"]:
        raise ValueError("aggregate document replacement counts differ from the manifest")
    for replacement in recipe.get("replacements") or []:
        old = replacement["old"]
        if aggregate_counts[old] < replacement.get("min_document_occurrences", 0):
            raise ValueError(f"document replacement minimum no longer holds for {old!r}")
        if expected_task_counts[old] < replacement.get("min_task_occurrences", 0):
            raise ValueError(f"task replacement minimum no longer holds for {old!r}")
    expected_tree_files = {
        "task.json",
        "augmentation-recipe.json",
        "mutation-manifest.json",
        *{f"documents/{relative}" for relative in actual},
    }
    actual_tree_files = {
        path.relative_to(task_dir).as_posix()
        for path in task_dir.rglob("*")
        if path.is_file()
    }
    if expected_tree_files != actual_tree_files:
        raise ValueError("generated tree contains missing or unexpected files")
    if not git_clean(source_root):
        raise ValueError("source Harvey worktree changed during validation")
    if not quiet:
        print(
            f"generated Harvey task valid: {manifest['variant_id']} "
            f"({len(actual)} documents, source {manifest['source_commit'][:12]})"
        )


def check_generated_root(
    output_root: Path = DEFAULT_OUTPUT,
    recipes_root: Path = DEFAULT_RECIPES,
    source_root: Path = DEFAULT_SOURCE,
) -> None:
    output_root = output_root.resolve()
    recipes_root = recipes_root.resolve()
    source_root = source_root.resolve()
    if not recipes_root.is_dir() or recipes_root.is_symlink():
        raise ValueError(f"recipes directory is missing or unsafe: {recipes_root}")
    recipe_paths = sorted(recipes_root.glob("*.json"))
    if not recipe_paths:
        raise ValueError(f"no augmentation recipes found in {recipes_root}")

    expected: dict[Path, tuple[Path, dict[str, Any]]] = {}
    variants: set[str] = set()
    for recipe_path in recipe_paths:
        if recipe_path.is_symlink():
            raise ValueError(f"recipe is a symlink: {recipe_path}")
        recipe = json.loads(recipe_path.read_text(encoding="utf-8"))
        validate_recipe(recipe)
        if recipe["variant_id"] in variants:
            raise ValueError(f"duplicate recipe variant_id: {recipe['variant_id']}")
        variants.add(recipe["variant_id"])
        relative = safe_relative(recipe["output_task"], "output_task")
        task_dir = output_root / "tasks" / Path(*relative.parts)
        ensure_safe_descendant(task_dir, (output_root / "tasks").resolve(), "generated task")
        if task_dir in expected:
            raise ValueError(f"multiple recipes target generated task: {task_dir}")
        expected[task_dir] = (recipe_path, recipe)

    expected_dirs = sorted(expected, key=str)
    for index, parent in enumerate(expected_dirs):
        for child in expected_dirs[index + 1:]:
            if child.is_relative_to(parent):
                raise ValueError(f"generated task paths overlap: {parent} and {child}")

    tasks_root = output_root / "tasks"
    manifest_dirs = {
        path.parent.resolve() for path in tasks_root.rglob("mutation-manifest.json")
    } if tasks_root.is_dir() else set()
    task_dirs = {
        path.parent.resolve() for path in tasks_root.rglob("task.json")
    } if tasks_root.is_dir() else set()
    if manifest_dirs != set(expected_dirs) or task_dirs != set(expected_dirs):
        missing = sorted(str(path) for path in set(expected_dirs) - (manifest_dirs & task_dirs))
        orphaned = sorted(str(path) for path in (manifest_dirs | task_dirs) - set(expected_dirs))
        raise ValueError(f"generated/recipe task set mismatch; missing={missing}, orphaned={orphaned}")

    for task_dir in expected_dirs:
        recipe_path, recipe = expected[task_dir]
        copied_recipe = task_dir / "augmentation-recipe.json"
        if copied_recipe.read_bytes() != canonical_json(recipe):
            raise ValueError(f"generated recipe copy differs from {recipe_path.name}")
        check_generated(task_dir, source_root, quiet=True)

    loose_files = [
        path for path in output_root.rglob("*")
        if path.is_file() and not any(path.is_relative_to(task_dir) for task_dir in expected_dirs)
    ] if output_root.is_dir() else []
    if loose_files:
        raise ValueError(
            "generated root contains files outside recipe task directories: "
            + ", ".join(str(path.relative_to(output_root)) for path in loose_files)
        )
    print(
        f"generated Harvey augmentation tree valid: {len(expected_dirs)} tasks, "
        f"{len(recipe_paths)} recipes, source {git_commit(source_root)[:12]}"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--recipe", type=Path)
    mode.add_argument("--check", type=Path)
    mode.add_argument("--check-root", type=Path)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--recipes-dir", type=Path, default=DEFAULT_RECIPES)
    parser.add_argument(
        "--replace-generated",
        action="store_true",
        help="replace only an existing generated task matching the recipe identity",
    )
    args = parser.parse_args()
    if args.check:
        if args.replace_generated:
            parser.error("--replace-generated requires --recipe")
        check_generated(args.check.resolve(), args.source.resolve())
        return 0
    if args.check_root:
        if args.replace_generated:
            parser.error("--replace-generated requires --recipe")
        check_generated_root(args.check_root, args.recipes_dir, args.source)
        return 0
    output = generate(
        args.recipe.resolve(),
        args.source.resolve(),
        args.output_root.resolve(),
        args.replace_generated,
    )
    print(f"generated Harvey task: {output}")
    check_generated(output, args.source.resolve())
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (
        FileExistsError,
        KeyError,
        OSError,
        RuntimeError,
        subprocess.CalledProcessError,
        TypeError,
        ValueError,
        zipfile.BadZipFile,
    ) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
